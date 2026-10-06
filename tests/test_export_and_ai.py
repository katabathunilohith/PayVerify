import io
import json
from datetime import datetime, timedelta
from types import SimpleNamespace

from openpyxl import load_workbook

from conftest import make_owner, png_bytes, receipt, upload, user_by_email
from payverify.extensions import db
from payverify.models import Payment
from payverify.services import extraction


def test_excel_has_one_tab_per_month(app, fake_ai):
    client = make_owner(app, "owner@example.com")
    fake_ai.data = receipt(utr="611111111111")
    upload(client, name="September Customer")
    fake_ai.data = receipt(utr="622222222222")
    upload(client, name="October Customer")
    with app.app_context():
        first = Payment.query.filter_by(utr="611111111111").one()
        first.received_at = datetime(2026, 9, 20, 6, 0)
        first.month_key = "2026-09"
        second = Payment.query.filter_by(utr="622222222222").one()
        second.received_at = datetime(2026, 10, 2, 6, 0)
        second.month_key = "2026-10"
        db.session.commit()
        second_id = second.id

    response = client.get("/export.xlsx")
    assert response.headers["Content-Disposition"].startswith("attachment")
    workbook = load_workbook(io.BytesIO(response.data))
    assert workbook.sheetnames == ["Sep 2026", "Oct 2026"]
    sheet = workbook["Oct 2026"]
    assert [cell.value for cell in sheet[1]][:5] == ["Date received", "Customer name", "Customer phone", "Amount", "UTR"]
    row = [cell.value for cell in sheet[2]]
    assert row[1] == "October Customer" and row[3] == 500 and row[4] == "622222222222" and row[7] == "Needs review"
    assert sheet.cell(row=2, column=10).hyperlink.target.endswith(f"/payments/{second_id}")

    single = load_workbook(io.BytesIO(client.get("/export.xlsx?month=2026-09").data))
    assert single.sheetnames == ["Sep 2026"]


class FakeResponse:
    def __init__(self, status_code=200, body=None):
        self.status_code = status_code
        self._body = body if body is not None else {}
        self.text = json.dumps(self._body)

    def json(self):
        return self._body


def _choice(content, finish_reason="stop"):
    return FakeResponse(200, {"choices": [{"message": {"role": "assistant", "content": content}, "finish_reason": finish_reason}]})


class FakePost:
    """Stands in for requests.post and records what would be sent to Meta."""

    def __init__(self, *responses):
        self.responses = list(responses)
        self.calls = []

    def __call__(self, url, headers=None, json=None, timeout=None):
        self.calls.append({"url": url, "headers": headers, "json": json, "timeout": timeout})
        return self.responses.pop(0)


def _read(app):
    with app.app_context():
        return extraction.extract_payment_details(png_bytes(), "image/png", datetime(2026, 10, 6).date(), "Asia/Kolkata")


def test_meta_request_shape(app, monkeypatch):
    app.config["META_MODEL_API_KEY"] = "test-key"
    fake = FakePost(_choice(json.dumps(receipt())))
    monkeypatch.setattr(extraction.requests, "post", fake)
    result = _read(app)
    assert result.ok and result.data["utr"] == "627812345678"
    [call] = fake.calls
    assert call["url"] == "https://api.meta.ai/v1/chat/completions"
    assert call["headers"]["Authorization"] == "Bearer test-key"
    body = call["json"]
    assert body["model"] == "muse-spark-1.3"
    assert body["reasoning_effort"] == "low"
    assert body["response_format"]["type"] == "json_schema"
    assert body["response_format"]["json_schema"]["schema"]["type"] == "object"
    parts = body["messages"][1]["content"]
    assert parts[0]["type"] == "text" and "2026-10-06" in parts[0]["text"]
    assert parts[1]["type"] == "image_url"
    assert parts[1]["image_url"]["url"].startswith("data:image/png;base64,")
    assert call["timeout"] <= 60


def test_meta_failures_are_reported_not_raised(app, monkeypatch):
    monkeypatch.setattr(extraction.time, "sleep", lambda _seconds: None)
    app.config["META_MODEL_API_KEY"] = ""
    missing = FakePost()
    monkeypatch.setattr(extraction.requests, "post", missing)
    result = _read(app)
    assert not result.ok and "META_MODEL_API_KEY" in result.error and missing.calls == []

    app.config["META_MODEL_API_KEY"] = "test-key"
    cases = [
        (FakePost(_choice("", finish_reason="refusal")), "declined", 1),
        (FakePost(FakeResponse(401, {"error": {"message": "bad key"}})), "rejected the API key", 1),
        (FakePost(FakeResponse(429), FakeResponse(429)), "busy", 2),  # retried once
        (FakePost(_choice("Sorry, I can't tell.")), "unreadable answer", 1),
    ]
    for fake, expected, calls in cases:
        monkeypatch.setattr(extraction.requests, "post", fake)
        result = _read(app)
        assert not result.ok and expected in result.error
        assert len(fake.calls) == calls

    # A busy server that recovers on the retry still works.
    recovering = FakePost(FakeResponse(503), _choice(json.dumps(receipt())))
    monkeypatch.setattr(extraction.requests, "post", recovering)
    assert _read(app).ok


def test_json_answers_in_fences_or_prose_are_accepted():
    assert extraction.parse_json_object('```json\n{"a": 1}\n```') == {"a": 1}
    assert extraction.parse_json_object('Here you go: {"a": 1} Done.') == {"a": 1}
    assert extraction.parse_json_object("[1, 2]") is None
    assert extraction.parse_json_object("nothing") is None


def test_dashboard_month_navigation_and_search(app, fake_ai):
    client = make_owner(app, "owner@example.com")
    upload(client, name="Findable Person", phone="919811111111")
    page = client.get("/?q=Findable").get_data(as_text=True)
    assert "Findable Person" in page
    assert "Findable Person" not in client.get("/?q=nobody").get_data(as_text=True)
    with app.app_context():
        month = Payment.query.one().month_key
        user = user_by_email("owner@example.com")
        assert user.timezone == "Asia/Kolkata"
    previous = (datetime.strptime(month + "-01", "%Y-%m-%d") - timedelta(days=1)).strftime("%Y-%m")
    assert "Findable Person" not in client.get(f"/?month={previous}").get_data(as_text=True)
    assert client.get("/?month=garbage").status_code == 200


def test_excel_does_not_run_customer_supplied_formulas(app, fake_ai):
    client = make_owner(app, "owner@example.com")
    upload(client, name='=HYPERLINK("http://evil.example","Click")')
    sheet = load_workbook(io.BytesIO(client.get("/export.xlsx").data)).worksheets[0]
    cell = sheet.cell(row=2, column=2)
    assert cell.data_type == "s"
    assert cell.value.startswith("=HYPERLINK")
