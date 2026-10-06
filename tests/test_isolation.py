"""Two accounts on one server must never see or touch each other's data."""

import io

from openpyxl import load_workbook

from conftest import make_owner, receipt, upload, user_by_email
from payverify.extensions import db
from payverify.models import BankCredit, Payment


def _payment_id(app, email):
    with app.app_context():
        return Payment.query.filter_by(user_id=user_by_email(email).id).one().id


def test_other_account_cannot_see_or_change_a_payment(app, fake_ai):
    alice = make_owner(app, "alice@example.com")
    bob = make_owner(app, "bob@example.com")
    upload(alice, name="Alice Customer")
    pid = _payment_id(app, "alice@example.com")

    assert alice.get(f"/payments/{pid}").status_code == 200
    assert alice.get(f"/payments/{pid}/screenshot").status_code == 200

    assert bob.get(f"/payments/{pid}").status_code == 404
    assert bob.get(f"/payments/{pid}/screenshot").status_code == 404
    assert bob.post(f"/payments/{pid}/status", data={"action": "verify"}).status_code == 404
    assert bob.post(f"/payments/{pid}/edit", data={"amount": "1"}).status_code == 404
    assert bob.post(f"/payments/{pid}/reread").status_code == 404
    assert bob.post(f"/payments/{pid}/delete").status_code == 404
    assert "Alice Customer" not in bob.get("/").get_data(as_text=True)
    assert "Alice Customer" not in bob.get("/?q=Alice").get_data(as_text=True)

    with app.app_context():
        payment = db.session.get(Payment, pid)
        assert payment.status != "verified"


def test_export_contains_only_own_payments(app, fake_ai):
    alice = make_owner(app, "alice@example.com")
    bob = make_owner(app, "bob@example.com")
    fake_ai.data = receipt(utr="611111111111")
    upload(alice, name="Alice Customer")
    fake_ai.data = receipt(utr="622222222222")
    upload(bob, name="Bob Customer")

    workbook = load_workbook(io.BytesIO(bob.get("/export.xlsx").data))
    cells = {str(cell.value) for sheet in workbook for row in sheet.iter_rows() for cell in row}
    assert "622222222222" in cells
    assert "611111111111" not in cells
    assert "Alice Customer" not in cells


def test_bank_credits_are_per_account(app, fake_ai):
    alice = make_owner(app, "alice@example.com")
    bob = make_owner(app, "bob@example.com")
    upload(alice)  # UTR 627812345678, ₹500

    # Bob receives a credit with the same UTR: Alice's payment must not verify.
    bob.post("/bank/manual", data={"amount": "500", "utr": "627812345678"})
    pid = _payment_id(app, "alice@example.com")
    with app.app_context():
        assert db.session.get(Payment, pid).status == "needs_review"
        credit_id = BankCredit.query.filter_by(user_id=user_by_email("bob@example.com").id).one().id
    assert alice.post(f"/bank/{credit_id}/delete").status_code == 404
    assert "627812345678" not in alice.get("/bank/").get_data(as_text=True)


def test_duplicate_check_is_per_account(app, fake_ai):
    alice = make_owner(app, "alice@example.com")
    bob = make_owner(app, "bob@example.com")
    upload(alice)
    upload(bob)  # same UTR, different business: not a duplicate for Bob
    with app.app_context():
        bob_payment = Payment.query.filter_by(user_id=user_by_email("bob@example.com").id).one()
        assert bob_payment.status == "needs_review"


def test_screenshots_stored_in_separate_folders(app, fake_ai):
    alice = make_owner(app, "alice@example.com")
    bob = make_owner(app, "bob@example.com")
    upload(alice)
    upload(bob)
    with app.app_context():
        alice_id = user_by_email("alice@example.com").id
        bob_id = user_by_email("bob@example.com").id
        paths = {p.user_id: p.screenshot_path for p in Payment.query.all()}
    assert paths[alice_id].startswith(f"{alice_id}/")
    assert paths[bob_id].startswith(f"{bob_id}/")
