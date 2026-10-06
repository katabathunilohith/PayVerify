"""Automatic approval mode, WhatsApp connection status, and database upgrades."""

import io

from openpyxl import load_workbook
from sqlalchemy import inspect, text

from conftest import make_owner, receipt, upload, user_by_email
from payverify import _add_missing_columns
from payverify.extensions import db
from payverify.models import Payment
from payverify.services.whatsapp import WhatsAppError

SMS = "Rs.{amount}.00 credited to a/c XX1234 (UPI Ref No {utr})."


def set_auto(client, on=True, next_url=None):
    data = {"mode": "automatic" if on else "bank"}
    if next_url:
        data["next"] = next_url
    return client.post("/settings/approval", data=data)


def only_payment(app, email="owner@example.com"):
    with app.app_context():
        p = Payment.query.filter_by(user_id=user_by_email(email).id).order_by(Payment.id.desc()).first()
        return {"status": p.status, "by": p.verified_by, "label": p.status_label, "reason": p.status_reason}


def test_auto_mode_approves_a_clean_screenshot(app, fake_ai):
    client = make_owner(app, "owner@example.com")
    set_auto(client)
    upload(client)
    payment = only_payment(app)
    assert payment["status"] == "verified"
    assert payment["by"] == "auto"
    assert payment["label"] == "Auto-approved"


def test_auto_mode_still_blocks_risky_screenshots(app, fake_ai):
    client = make_owner(app, "owner@example.com")
    set_auto(client)
    cases = [
        ({**receipt(utr="600000000011"), "looks_edited": True}, "edited"),
        (receipt(utr="600000000012", status="unknown"), "doesn't clearly say"),
        (receipt(utr="600000000013", payee="someone.else@ybl"), "isn't one of your UPI IDs"),
        (receipt(utr="600000000014", minutes_ago=60 * 24 * 9), "old screenshot"),
    ]
    for data, expected in cases:
        fake_ai.data = data
        upload(client)
        payment = only_payment(app)
        assert payment["status"] == "needs_review", expected
        assert expected in payment["reason"]

    fake_ai.data = receipt(utr="600000000014", minutes_ago=5)
    upload(client)
    assert only_payment(app)["status"] == "duplicate"


def test_auto_mode_needs_upi_ids(app, fake_ai):
    client = make_owner(app, "owner@example.com", upi="")
    set_auto(client)
    upload(client)
    payment = only_payment(app)
    assert payment["status"] == "needs_review"
    assert "add your UPI IDs" in payment["reason"]


def test_bank_sms_confirms_or_overrules_auto_approval(app, fake_ai):
    client = make_owner(app, "owner@example.com")
    set_auto(client)
    with app.app_context():
        sms_url = f"/webhooks/sms/{user_by_email('owner@example.com').sms_token}"

    fake_ai.data = receipt(utr="600000000021", amount=500)
    upload(client)
    client.post(sms_url, json={"from": "AX-HDFCBK", "text": SMS.format(amount=500, utr="600000000021")})
    confirmed = only_payment(app)
    assert confirmed["status"] == "verified" and confirmed["by"] == "bank"

    fake_ai.data = receipt(utr="600000000022", amount=5000)
    upload(client)
    assert only_payment(app)["by"] == "auto"
    client.post(sms_url, json={"from": "AX-HDFCBK", "text": SMS.format(amount=500, utr="600000000022")})
    overruled = only_payment(app)
    assert overruled["status"] == "needs_review"
    assert "bank received ₹500" in overruled["reason"]


def test_turning_auto_on_approves_waiting_payments_and_off_keeps_them(app, fake_ai):
    client = make_owner(app, "owner@example.com")
    upload(client)
    assert only_payment(app)["status"] == "needs_review"

    response = set_auto(client, on=True, next_url="/?month=2026-10")
    assert response.headers["Location"] == "/?month=2026-10"
    assert only_payment(app)["status"] == "verified"

    set_auto(client, on=False)
    with app.app_context():
        assert not user_by_email("owner@example.com").auto_approve
    assert only_payment(app)["status"] == "verified"  # not taken back

    fake_ai.data = receipt(utr="600000000031")
    upload(client)
    assert only_payment(app)["status"] == "needs_review"


def test_auto_toggle_rejects_bad_input_and_foreign_redirects(app):
    client = make_owner(app, "owner@example.com")
    assert client.post("/settings/approval", data={"mode": "yolo"}).status_code == 400
    response = set_auto(client, next_url="https://evil.example/")
    assert response.headers["Location"] == "/settings/"


def test_excel_total_counts_auto_approved(app, fake_ai):
    client = make_owner(app, "owner@example.com")
    set_auto(client)
    upload(client)
    sheet = load_workbook(io.BytesIO(client.get("/export.xlsx").data)).worksheets[0]
    assert sheet.cell(row=2, column=8).value == "Auto-approved"
    assert '"Auto-approved"' in sheet.cell(row=4, column=4).value


def test_dashboard_shows_the_switch(app):
    client = make_owner(app, "owner@example.com")
    page = client.get("/").get_data(as_text=True)
    assert 'role="switch" aria-checked="false"' in page
    set_auto(client)
    assert 'role="switch" aria-checked="true"' in client.get("/").get_data(as_text=True)


def test_whatsapp_status_checklist(app, monkeypatch):
    client = make_owner(app, "owner@example.com")
    with app.app_context():
        user = user_by_email("owner@example.com")
        user.wa_phone_number_id = "1065"
        user.wa_access_token = "token"
        user.wa_app_secret = "secret"
        db.session.commit()
        token, verify = user.webhook_token, user.wa_verify_token

    page = client.get("/settings/whatsapp").get_data(as_text=True)
    assert "Not yet. Paste the callback URL" in page

    client.get(f"/webhooks/whatsapp/{token}?hub.mode=subscribe&hub.verify_token={verify}&hub.challenge=7")
    assert "Meta checked your callback URL" in client.get("/settings/whatsapp").get_data(as_text=True)

    sent = []
    monkeypatch.setattr("payverify.services.whatsapp.send_test_template", lambda user, to: sent.append(to))
    client.post("/settings/whatsapp", data={"action": "send_test", "test_phone": "+91 98765 43210"})
    assert sent == ["919876543210"]

    client.post("/settings/whatsapp", data={"action": "rotate"})
    with app.app_context():
        assert user_by_email("owner@example.com").wa_verified_at is None


def test_whatsapp_errors_explain_the_fix():
    error = WhatsAppError("Recipient phone number not in allowed list (code 131030).", 131030)
    assert "API Setup" in str(error)
    assert str(WhatsAppError("Something odd.", 999)) == "Something odd."


def test_old_database_gets_new_columns(app):
    make_owner(app, "owner@example.com")
    with app.app_context():
        with db.engine.begin() as connection:
            connection.execute(text('ALTER TABLE users DROP COLUMN approval_mode'))
        db.session.remove()
        assert "approval_mode" not in {c["name"] for c in inspect(db.engine).get_columns("users")}
        _add_missing_columns()
        assert user_by_email("owner@example.com").approval_mode == "bank"


def test_check_command_reports_status(app, monkeypatch):
    make_owner(app, "owner@example.com")
    runner = app.test_cli_runner()

    app.config["META_MODEL_API_KEY"] = ""
    result = runner.invoke(args=["check"])
    assert result.exit_code == 0
    assert "✘ Meta AI key: not set" in result.output
    assert "✔ Database: connected" in result.output and "1 account(s)" in result.output
    assert "Account owner@example.com" in result.output
    assert "✘ Auto-approve: off" in result.output

    from payverify.services import extraction

    class Ok:
        status_code = 200
        text = "{}"

    class Rejected:
        status_code = 401
        text = "{}"

    app.config["META_MODEL_API_KEY"] = "test-key"
    monkeypatch.setattr(extraction.requests, "post", lambda *a, **k: Ok())
    assert "✔ Meta AI key: works" in runner.invoke(args=["check"]).output
    monkeypatch.setattr(extraction.requests, "post", lambda *a, **k: Rejected())
    assert "Meta rejected the key" in runner.invoke(args=["check"]).output
