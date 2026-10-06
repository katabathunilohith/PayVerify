"""Import credits from a bank statement (CSV or XLSX).

Every bank lays its statement out differently, so this finds the heading row
by its column names (date, narration, reference, credit/deposit) and reads
the UPI reference from the reference column or the narration.
"""

import csv
import io
import os
import re
from datetime import date, datetime

from openpyxl import load_workbook

from ..utils import parse_amount
from .sms_parser import TWELVE_DIGITS_RE, ParsedCredit

DATE_FORMATS = (
    "%d/%m/%Y", "%d/%m/%y", "%d-%m-%Y", "%d-%m-%y", "%d.%m.%Y", "%d.%m.%y",
    "%d %b %Y", "%d-%b-%Y", "%d-%b-%y", "%d %b %y", "%d %B %Y", "%Y-%m-%d",
    "%d/%m/%Y %H:%M:%S", "%d-%m-%Y %H:%M:%S", "%Y-%m-%d %H:%M:%S", "%d/%m/%Y %H:%M", "%d-%m-%Y %H:%M",
)
NARRATION_WORDS = ("narration", "description", "particulars", "remarks", "details")
REFERENCE_WORDS = ("ref", "chq", "cheque", "utr", "rrn")
TYPE_HEADINGS = {"dr/cr", "cr/dr", "type", "transaction type", "debit/credit", "dr / cr", "cr / dr"}


class StatementError(ValueError):
    pass


def _cell_text(value):
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def read_rows(filename, data):
    ext = os.path.splitext(filename or "")[1].lower()
    if ext in (".csv", ".txt"):
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = data.decode("latin-1")
        try:
            dialect = csv.Sniffer().sniff(text[:5000], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        return [[cell.strip() for cell in row] for row in csv.reader(io.StringIO(text), dialect)]
    if ext == ".xlsx":
        try:
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as exc:
            raise StatementError("That Excel file couldn't be opened.") from exc
        sheet = workbook.worksheets[0]
        rows = [[_cell_text(value) for value in row] for row in sheet.iter_rows(values_only=True)]
        workbook.close()
        return rows
    raise StatementError("Upload a CSV or XLSX file. If your bank gives an .xls file, open it and save it as .xlsx or CSV.")


def _first(cells, test):
    return next((i for i, cell in enumerate(cells) if cell and test(cell)), None)


def find_columns(rows):
    for index, row in enumerate(rows[:60]):
        cells = [" ".join(cell.lower().split()) for cell in row]
        columns = {
            "date": _first(cells, lambda c: "date" in c or c in ("dt", "tran dt", "txn dt")),
            "credit": _first(cells, lambda c: ("credit" in c or "deposit" in c or c in ("cr", "cr.", "cr amount"))
                             and "debit" not in c and "withdraw" not in c and "bal" not in c and "/" not in c),
            "amount": _first(cells, lambda c: "amount" in c and "debit" not in c and "withdraw" not in c),
            "type": _first(cells, lambda c: c in TYPE_HEADINGS),
            "narration": _first(cells, lambda c: any(word in c for word in NARRATION_WORDS)),
            "reference": _first(cells, lambda c: any(word in c for word in REFERENCE_WORDS)),
        }
        has_amount = columns["credit"] is not None or (columns["amount"] is not None and columns["type"] is not None)
        if columns["date"] is not None and has_amount:
            return index, columns
    raise StatementError(
        "Couldn't find the column headings. The file needs a date column and a credit/deposit column."
    )


def parse_date(text):
    text = " ".join((text or "").split())
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    return None


def _get(row, column):
    return row[column] if column is not None and column < len(row) else ""


def parse_statement(rows):
    """Return (credits, skipped_without_utr)."""
    header, columns = find_columns(rows)
    credits, skipped = [], 0
    for row in rows[header + 1:]:
        if not any(row):
            continue
        if columns["credit"] is not None:
            amount = parse_amount(_get(row, columns["credit"]))
        else:
            kind = _get(row, columns["type"]).strip().lower()
            amount = parse_amount(_get(row, columns["amount"])) if kind.startswith("c") else None
        if not amount:
            continue
        narration = _get(row, columns["narration"])
        reference = re.sub(r"\D", "", _get(row, columns["reference"]))
        if len(reference) > 12 and not reference[:-12].strip("0"):
            reference = reference[-12:]  # HDFC pads references with zeros: 0000427812345678
        if len(reference) == 12:
            utr = reference
        else:
            match = TWELVE_DIGITS_RE.search(narration) or TWELVE_DIGITS_RE.search(" ".join(row))
            utr = match.group(1) if match else None
        if not utr:
            skipped += 1
            continue
        credits.append(ParsedCredit(amount, utr, parse_date(_get(row, columns["date"])), narration[:120]))
    return credits, skipped
