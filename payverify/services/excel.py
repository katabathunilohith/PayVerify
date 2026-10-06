"""Excel export: one workbook, one tab per month ("Oct 2026")."""

import io
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..utils import local_now, month_key, month_label, to_local

HEADERS = [
    ("Date received", 18), ("Customer name", 22), ("Customer phone", 16), ("Amount", 13),
    ("UTR", 16), ("Payment app", 14), ("Payment date & time", 20), ("Status", 14),
    ("Reason / notes", 48), ("Screenshot", 12),
]
# Indian digit grouping (12,50,000.00)
INR_FORMAT = '[>=10000000]"₹"##\\,##\\,##\\,##0.00;[>=100000]"₹"##\\,##\\,##0.00;"₹"##,##0.00'
DATE_FORMAT = "dd-mmm-yyyy hh:mm AM/PM"
STATUS_FILLS = {
    "verified": "D1F2E2", "needs_review": "FDEFD2", "duplicate": "FAD4D4",
    "rejected": "E8E8EC", "processing": "DCE6FB",
}
HEADER_FILL = PatternFill("solid", fgColor="312E81")


def _sheet(workbook, key, payments, tz_name, link_for):
    ws = workbook.create_sheet(title=month_label(key))
    for col, (title, width) in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A2"

    row = 1
    for row, payment in enumerate(payments, start=2):
        notes = payment.status_reason
        if payment.notes:
            notes = f"{notes} | Note: {payment.notes}" if notes else payment.notes
        values = [
            to_local(payment.received_at, tz_name),
            payment.customer_name,
            payment.customer_phone,
            payment.amount_paise / 100 if payment.amount_paise else None,
            payment.utr or "",
            payment.payment_app,
            payment.paid_at,
            payment.status_label,
            notes,
            "Open" if payment.screenshot_path else "",
        ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col, value=value)
            if isinstance(value, str) and value.startswith("="):
                # Customer names come from WhatsApp profiles: never let them run as formulas.
                cell.data_type = "s"
        ws.cell(row=row, column=1).number_format = DATE_FORMAT
        ws.cell(row=row, column=4).number_format = INR_FORMAT
        ws.cell(row=row, column=5).number_format = "@"
        ws.cell(row=row, column=7).number_format = DATE_FORMAT
        ws.cell(row=row, column=8).fill = PatternFill("solid", fgColor=STATUS_FILLS.get(payment.status, "FFFFFF"))
        if payment.screenshot_path:
            link = ws.cell(row=row, column=10)
            link.hyperlink = link_for(payment)
            link.style = "Hyperlink"

    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{max(row, 1)}"
    total_row = row + 2
    ws.cell(row=total_row, column=3, value="Verified total").font = Font(bold=True)
    last = max(row, 2)
    total = ws.cell(row=total_row, column=4,
                    value=f'=SUMIFS(D2:D{last},H2:H{last},"Verified")+SUMIFS(D2:D{last},H2:H{last},"Auto-approved")')
    total.number_format = INR_FORMAT
    total.font = Font(bold=True)


def build_workbook(user, payments, link_for):
    """payments: this account's payments. link_for(payment) -> screenshot URL."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    by_month = defaultdict(list)
    for payment in sorted(payments, key=lambda p: p.received_at):
        by_month[payment.month_key].append(payment)
    if not by_month:
        by_month[month_key(local_now(user.timezone))] = []
    for key in sorted(by_month):
        _sheet(workbook, key, by_month[key], user.timezone, link_for)
    out = io.BytesIO()
    workbook.save(out)
    return out.getvalue()
